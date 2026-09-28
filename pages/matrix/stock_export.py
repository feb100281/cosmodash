# pages/matrix/stock_export.py

from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any, Iterable

import pandas as pd
from datetime import datetime

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .xl_brand import (
    EXPENSE,
    FMT_QTY,
    OCCUPIED_BG,
    WARN,
    WARN_BG,
    build_toc,
    fill,
    finalize,
    font,
    set_number_format,
    style_table_sheet,
)


UNALLOCATED_LABEL = "НЕ РАСПРЕДЕЛЕНО В ИСТОЧНИКЕ"
EPS = 1e-9


SUMMARY_COLUMNS = [
    ("cat_name", "Категория"),
    ("sc_name", "Подкатегория"),
    ("manu", "Производитель"),
    ("fullname", "Номенклатура"),
    ("article", "Артикул"),
    ("barcode", "Штрихкоды товара"),
    ("stock_available", "Доступный остаток"),
    ("stock_ordered", "Заказано"),
    ("stock_total", "Всего с заказами"),
    ("stock_status", "Статус запаса"),
    ("barcode_stocks_display", "Остатки по штрихкодам"),
]


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _to_float(value: Any) -> float:
    value = pd.to_numeric(
        pd.Series([value]),
        errors="coerce",
    ).fillna(0).iloc[0]
    return float(value)


def _parse_qty_list(value: Any) -> dict[str, float]:
    """
    Разбирает детализацию остатков вида:

        ["Европарк - 2 шт.", "ОСНОВНОЙ склад - 1364 шт."]

    или:

        ["2000000001944 - 23 шт.", "2000000001944 - 984 шт."]

    Повторяющиеся склады / штрихкоды суммируются.
    """
    if _is_empty(value):
        return {}

    raw_values = None

    if isinstance(value, (list, tuple)):
        raw_values = list(value)
    else:
        text = str(value).strip()

        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                raw_values = parsed
        except (json.JSONDecodeError, TypeError, ValueError):
            raw_values = None

    result: dict[str, float] = {}

    if raw_values is not None:
        for item in raw_values:
            match = re.match(
                r"^\s*(.+?)\s*-\s*(-?\d+(?:[.,]\d+)?)\s*шт\.?\s*$",
                str(item),
                flags=re.IGNORECASE,
            )

            if not match:
                continue

            name = match.group(1).strip().strip("'\"").strip()
            if not name:
                continue

            try:
                qty = float(match.group(2).replace(",", "."))
            except (TypeError, ValueError):
                continue

            result[name] = result.get(name, 0.0) + qty

        return result

    matches = re.findall(
        r"""["']?([^"'\[\],]+?)["']?\s*-\s*
            (-?\d+(?:[.,]\d+)?)\s*шт\.?""",
        str(value),
        flags=re.IGNORECASE | re.VERBOSE,
    )

    for name, qty in matches:
        name = str(name).strip().strip("'\"").strip()
        if not name:
            continue

        try:
            qty_value = float(str(qty).replace(",", "."))
        except (TypeError, ValueError):
            continue

        result[name] = result.get(name, 0.0) + qty_value

    return result


def _dynamic_columns(
    columns: Iterable[str],
    prefix: str,
) -> list[str]:
    return sorted(
        [
            str(column)
            for column in columns
            if str(column).startswith(prefix)
        ],
        key=lambda value: value.split("::", 1)[-1].lower(),
    )


def _base_row(row: pd.Series) -> dict[str, Any]:
    return {
        "Категория": row.get("cat_name", ""),
        "Подкатегория": row.get("sc_name", ""),
        "Производитель": row.get("manu", ""),
        "Номенклатура": row.get("fullname", ""),
        "Артикул": row.get("article", ""),
    }


def _prepare_summary_df(df_matrix: pd.DataFrame) -> pd.DataFrame:
    if df_matrix is None or df_matrix.empty:
        return pd.DataFrame(
            columns=[title for _, title in SUMMARY_COLUMNS]
        )

    df = df_matrix.copy()

    warehouse_columns = _dynamic_columns(df.columns, "stock_wh::")
    ordered_warehouse_columns = _dynamic_columns(df.columns, "ordered_wh::")

    base_columns = [
        source
        for source, _ in SUMMARY_COLUMNS
        if source in df.columns
    ]

    result = df[
        base_columns
        + warehouse_columns
        + ordered_warehouse_columns
    ].copy()

    rename_map = {
        source: title
        for source, title in SUMMARY_COLUMNS
        if source in result.columns
    }

    for column in warehouse_columns:
        warehouse = column.split("::", 1)[1]
        rename_map[column] = f"Магазин | {warehouse}"

    for column in ordered_warehouse_columns:
        warehouse = column.split("::", 1)[1]
        rename_map[column] = f"Заказ | {warehouse}"

    return result.rename(columns=rename_map)


def _warehouse_breakdown(
    row: pd.Series,
    warehouse_columns: list[str],
) -> tuple[list[tuple[str, float]], float]:
    values: list[tuple[str, float]] = []
    total = 0.0

    for column in warehouse_columns:
        qty = _to_float(row.get(column, 0))

        if abs(qty) <= EPS:
            continue

        warehouse = column.split("::", 1)[1]
        values.append((warehouse, qty))
        total += qty

    return values, total


def _barcode_breakdown(
    row: pd.Series,
) -> tuple[list[tuple[str, float]], float]:
    parsed = _parse_qty_list(
        row.get("barcode_stocks")
    )

    values: list[tuple[str, float]] = []
    total = 0.0

    for barcode, qty in sorted(
        parsed.items(),
        key=lambda item: str(item[0]),
    ):
        qty = float(qty)

        if abs(qty) <= EPS:
            continue

        values.append((str(barcode), qty))
        total += qty

    return values, total


def _prepare_warehouse_df(
    df_matrix: pd.DataFrame,
) -> pd.DataFrame:
    columns = [
        "Категория",
        "Подкатегория",
        "Производитель",
        "Номенклатура",
        "Артикул",
        "Магазин",
        "Остаток",
    ]

    if df_matrix is None or df_matrix.empty:
        return pd.DataFrame(columns=columns)

    warehouse_columns = _dynamic_columns(
        df_matrix.columns,
        "stock_wh::",
    )

    rows: list[dict[str, Any]] = []

    for _, row in df_matrix.iterrows():
        base = _base_row(row)

        available = _to_float(
            row.get("stock_available", 0)
        )

        breakdown, detailed_total = _warehouse_breakdown(
            row,
            warehouse_columns,
        )

        for warehouse, qty in breakdown:
            rows.append(
                {
                    **base,
                    "Магазин": warehouse,
                    "Остаток": qty,
                }
            )

        difference = available - detailed_total

        if abs(difference) > EPS:
            rows.append(
                {
                    **base,
                    "Магазин": UNALLOCATED_LABEL,
                    "Остаток": difference,
                }
            )

    return pd.DataFrame(rows, columns=columns)


def _prepare_barcode_df(
    df_matrix: pd.DataFrame,
) -> pd.DataFrame:
    columns = [
        "Категория",
        "Подкатегория",
        "Производитель",
        "Номенклатура",
        "Артикул",
        "Штрихкод",
        "Остаток",
    ]

    if df_matrix is None or df_matrix.empty:
        return pd.DataFrame(columns=columns)

    rows: list[dict[str, Any]] = []

    for _, row in df_matrix.iterrows():
        base = _base_row(row)

        available = _to_float(
            row.get("stock_available", 0)
        )

        breakdown, detailed_total = _barcode_breakdown(row)

        for barcode, qty in breakdown:
            rows.append(
                {
                    **base,
                    "Штрихкод": barcode,
                    "Остаток": qty,
                }
            )

        difference = available - detailed_total

        if abs(difference) > EPS:
            rows.append(
                {
                    **base,
                    "Штрихкод": UNALLOCATED_LABEL,
                    "Остаток": difference,
                }
            )

    return pd.DataFrame(rows, columns=columns)


def _prepare_control_df(
    df_matrix: pd.DataFrame,
) -> pd.DataFrame:
    columns = [
        "Категория",
        "Подкатегория",
        "Производитель",
        "Номенклатура",
        "Артикул",
        "Доступный остаток",
        "По магазинам",
        "Расхождение магазины",
        "По штрихкодам",
        "Расхождение штрихкоды",
    ]

    if df_matrix is None or df_matrix.empty:
        return pd.DataFrame(columns=columns)

    warehouse_columns = _dynamic_columns(
        df_matrix.columns,
        "stock_wh::",
    )

    rows: list[dict[str, Any]] = []

    for _, row in df_matrix.iterrows():
        available = _to_float(
            row.get("stock_available", 0)
        )

        _, warehouse_total = _warehouse_breakdown(
            row,
            warehouse_columns,
        )

        _, barcode_total = _barcode_breakdown(row)

        rows.append(
            {
                **_base_row(row),
                "Доступный остаток": available,
                "По магазинам": warehouse_total,
                "Расхождение магазины": (
                    available - warehouse_total
                ),
                "По штрихкодам": barcode_total,
                "Расхождение штрихкоды": (
                    available - barcode_total
                ),
            }
        )

    result = pd.DataFrame(rows, columns=columns)

    result["_problem_size"] = (
        result["Расхождение магазины"].abs()
        + result["Расхождение штрихкоды"].abs()
    )

    return (
        result
        .sort_values(
            "_problem_size",
            ascending=False,
        )
        .drop(columns="_problem_size")
        .reset_index(drop=True)
    )


QTY_HEADERS = [
    "Доступный остаток",
    "Заказано",
    "Всего с заказами",
    "Остаток",
    "По магазинам",
    "Расхождение магазины",
    "По штрихкодам",
    "Расхождение штрихкоды",
]

WIDTHS = {
    "Категория": 22,
    "Подкатегория": 24,
    "Производитель": 22,
    "Номенклатура": 46,
    "Артикул": 18,
    "Штрихкоды товара": 28,
    "Доступный остаток": 14,
    "Заказано": 12,
    "Всего с заказами": 14,
    "Статус запаса": 26,
    "Остатки по штрихкодам": 44,
    "Магазин": 32,
    "Штрихкод": 28,
    "Остаток": 12,
}


def _style_sheet(ws, *, title: str, stock_date: str | None = None) -> None:
    subtitle = f"Остатки на {stock_date}" if stock_date else "Текущие остатки"
    hr = style_table_sheet(
        ws,
        title=title,
        subtitle=subtitle,
        header_row=1,
        freeze_col=6 if ws.max_column >= 6 else 2,
        key_headers=["Доступный остаток", "Остаток"],
        wrap_headers=["Номенклатура", "Остатки по штрихкодам"],
    )
    set_number_format(ws, hr, QTY_HEADERS, FMT_QTY)
    set_number_format(ws, hr, ["Магазин | ", "Заказ | "], FMT_QTY, prefix=True)

    for cc in range(1, ws.max_column + 1):
        h = str(ws.cell(row=hr, column=cc).value or "")
        width = WIDTHS.get(h)
        if width is None and (h.startswith("Магазин | ") or h.startswith("Заказ | ")):
            width = 18
        ws.column_dimensions[get_column_letter(cc)].width = width or 16

    for r in range(hr + 1, ws.max_row + 1):
        values = [str(ws.cell(row=r, column=cc).value or "") for cc in range(1, ws.max_column + 1)]
        if any(UNALLOCATED_LABEL in v for v in values):
            for cc in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=cc)
                cell.fill = fill(WARN_BG)
                cell.font = font(10, True, WARN)

    for cc in range(1, ws.max_column + 1):
        if str(ws.cell(row=hr, column=cc).value or "").startswith("Расхождение"):
            for r in range(hr + 1, ws.max_row + 1):
                cell = ws.cell(row=r, column=cc)
                try:
                    bad = abs(float(cell.value or 0)) > EPS
                except (TypeError, ValueError):
                    bad = False
                if bad:
                    cell.fill = fill(OCCUPIED_BG)
                    cell.font = font(10, True, EXPENSE)


def _get_stock_date_label(
    df_matrix: pd.DataFrame,
) -> str | None:
    """
    Возвращает дату остатков для заголовка Excel.

    Если в выгрузке одна дата — показываем:
        24.07.2026

    Если по каким-то причинам дат несколько —
    показываем диапазон:
        23.07.2026–24.07.2026
    """
    if (
        df_matrix is None
        or df_matrix.empty
        or "stock_date" not in df_matrix.columns
    ):
        return None

    dates = (
        pd.to_datetime(
            df_matrix["stock_date"],
            errors="coerce",
        )
        .dropna()
        .dt.normalize()
        .drop_duplicates()
        .sort_values()
    )

    if dates.empty:
        return None

    first_date = dates.iloc[0].strftime("%d.%m.%Y")
    last_date = dates.iloc[-1].strftime("%d.%m.%Y")

    if first_date == last_date:
        return first_date

    return f"{first_date}–{last_date}"


def build_stock_excel_bytes(df_matrix: pd.DataFrame) -> bytes:
    """Остатки: оглавление, «Остатки», «По магазинам», «По штрихкодам»."""
    stock_date = _get_stock_date_label(df_matrix)

    summary_df = _prepare_summary_df(df_matrix)
    warehouse_df = _prepare_warehouse_df(df_matrix)
    barcode_df = _prepare_barcode_df(df_matrix)

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Остатки", index=False)
        warehouse_df.to_excel(writer, sheet_name="По магазинам", index=False)
        barcode_df.to_excel(writer, sheet_name="По штрихкодам", index=False)

        wb = writer.book
        _style_sheet(wb["Остатки"], title="ТЕКУЩИЕ ОСТАТКИ", stock_date=stock_date)
        _style_sheet(wb["По магазинам"], title="ОСТАТКИ ПО МАГАЗИНАМ", stock_date=stock_date)
        _style_sheet(wb["По штрихкодам"], title="ОСТАТКИ ПО ШТРИХКОДАМ", stock_date=stock_date)

        def _sum(col):
            if col not in summary_df.columns:
                return None
            return float(pd.to_numeric(summary_df[col], errors="coerce").fillna(0).sum())

        sku = int((pd.to_numeric(summary_df.get("Доступный остаток", 0), errors="coerce").fillna(0) > 0).sum()) \
            if "Доступный остаток" in summary_df.columns else len(summary_df)

        build_toc(
            wb,
            title="Остатки товаров",
            subtitle="Доступный остаток и заказанный товар по SKU, магазинам и штрихкодам",
            params=f"Штуки · остатки на {stock_date or '—'} · выгрузка {datetime.now():%d.%m.%Y}",
            cards=[
                ("ДОСТУПНО, ШТ.", _sum("Доступный остаток"), FMT_QTY, "весь доступный остаток"),
                ("ЗАКАЗАНО, ШТ.", _sum("Заказано"), FMT_QTY, "в пути / в заказе"),
                ("ВСЕГО С ЗАКАЗАМИ, ШТ.", _sum("Всего с заказами"), FMT_QTY, "доступно + заказано"),
                ("SKU С ОСТАТКОМ", sku, FMT_QTY, f"из {len(summary_df)} SKU"),
            ],
            sheets=[
                ("Остатки", "Одна строка — один SKU: остаток, заказ, статус запаса"),
                ("По магазинам", "Остаток каждого SKU в разрезе магазинов и складов"),
                ("По штрихкодам", "Остаток каждого SKU в разрезе штрихкодов"),
            ],
        )
        finalize(wb, ["Оглавление", "Остатки", "По магазинам", "По штрихкодам"])

    output.seek(0)
    return output.getvalue()
