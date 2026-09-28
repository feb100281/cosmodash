# pages/matrix/turnover_export.py
"""Excel-выгрузка анализа оборачиваемости."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .turnover import (
    abc_turnover_matrix,
    STATUS_DEAD,
    STATUS_VERY_SLOW,
    build_turnover_insights,
    turnover_by_category,
    turnover_by_warehouse,
    turnover_by_status,
    turnover_kpis,
)

NAVY = "2F6656"
NAVY_2 = "3D7A67"
NAVY_3 = "1F5E4E"
TEXT = "1F1F1F"
TEXT_2 = "4A4A4A"
MUTED = "8A8A8A"
SURFACE_2 = "F3F8F6"
SURFACE_3 = "EDF5F1"
SURFACE_4 = "E7F1ED"
TOTAL_ROW = "E7F1ED"
ZEBRA_ROW = "F7F7F7"
LINE = "D9D9D9"
EXPENSE = "7B4437"
INFO = "2F75B5"
INFO_BG = "EAF2FB"
WARN = "9A6100"
WARN_BG = "FDF3DE"
OCCUPIED_BG = "F6E9E4"

FONT = "Roboto"

FMT_MONEY = '#,##0;(#,##0);"–"'
FMT_PCT = '#,##0.0" %";(#,##0.0" %");"–"'
FMT_QTY = '#,##0;(#,##0);"–"'
FMT_DEC = '#,##0.00;(#,##0.00);"–"'
FMT_DATE = "dd.mm.yyyy"

TOC_SHEET_NAME = "Оглавление"


def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def _side(color=LINE, style="thin"):
    return Side(style=style, color=color)


def _font(size=10, bold=False, color=TEXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=None)


def _write(ws, r, c, v, *, font=None, fill=None, fmt=None, align="left", wrap=False, border=None, indent=0):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font or _font()
    if fill:
        cell.fill = _fill(fill)
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap, indent=indent)
    if border:
        cell.border = border
    return cell


def _setup(ws, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True


def _back_to_toc(ws, width):
    width = max(width, 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value="←  Оглавление")
    c.hyperlink = f"#'{TOC_SHEET_NAME}'!A1"
    c.font = Font(name=FONT, size=9, bold=True, color=NAVY_3, underline=None)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for cc in range(1, width + 1):
        cell = ws.cell(row=1, column=cc)
        cell.fill = _fill(SURFACE_4)
        cell.border = Border(bottom=_side())
    ws.row_dimensions[1].height = 16


def _sheet_header(ws, width, title, subtitle, params):
    width = max(width, 2)
    _back_to_toc(ws, width)
    _write(ws, 2, 1, title, font=_font(14, True))
    ws.row_dimensions[2].height = 24
    _write(ws, 3, 1, subtitle, font=_font(9, color=TEXT_2))
    ws.row_dimensions[3].height = 14
    _write(ws, 4, 1, params, font=_font(9, color=MUTED))
    ws.row_dimensions[4].height = 14
    for cc in range(1, width + 1):
        ws.cell(row=5, column=cc).border = Border(bottom=_side(NAVY, "medium"))
    ws.row_dimensions[5].height = 6
    return 6


def _table(ws, start_row, columns, rows, *, total=None, widths=None, autofilter=False, freeze=True):
    """columns: [(title, fmt, align)]; total: строка итогов или None."""
    n = len(columns)
    hdr = start_row
    for j, (title, _, align) in enumerate(columns, start=1):
        c = _write(ws, hdr, j, title, font=_font(10, True, "FFFFFF"), fill=NAVY,
                   align="center" if align != "left" else "left", wrap=True,
                   border=Border(left=_side(), right=_side(), top=_side(), bottom=_side()))
    ws.row_dimensions[hdr].height = 30

    r = hdr
    for row in rows:
        r += 1
        for j, ((_, fmt, align), v) in enumerate(zip(columns, row), start=1):
            if isinstance(v, float) and np.isnan(v):
                v = None
            _write(ws, r, j, v, fmt=fmt, align=align, border=Border(bottom=_side()))
    last = r

    # зебра по чётным строкам без заливки
    for rr in range(hdr + 1, last + 1):
        if rr % 2:
            continue
        for cc in range(1, n + 1):
            cell = ws.cell(row=rr, column=cc)
            if not cell.fill.fill_type:
                cell.fill = _fill(ZEBRA_ROW)

    if total is not None:
        r += 1
        for j, ((_, fmt, align), v) in enumerate(zip(columns, total), start=1):
            if isinstance(v, float) and np.isnan(v):
                v = None
            _write(ws, r, j, v, font=_font(10, True), fill=TOTAL_ROW, fmt=fmt, align=align,
                   border=Border(top=_side(NAVY, "medium"), bottom=_side(NAVY, "double")))

    # разделители колонок
    for rr in range(hdr + 1, r + 1):
        for cc in range(1, n):
            cell = ws.cell(row=rr, column=cc)
            old = cell.border
            cell.border = Border(left=old.left, right=_side(), top=old.top, bottom=old.bottom)

    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    if freeze:
        ws.freeze_panes = ws.cell(row=hdr + 1, column=2)
    if autofilter and last > hdr:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(n)}{last}"
    return r


def _to_date(v):
    if v in (None, "") or (isinstance(v, float) and np.isnan(v)):
        return None
    try:
        return datetime.strptime(str(v), "%d.%m.%Y")
    except ValueError:
        return None


def _n(v):
    try:
        f = float(v)
        return None if np.isnan(f) else f
    except (TypeError, ValueError):
        return None


def build_turnover_excel_bytes(df: pd.DataFrame, period_label: str = "", scope: str = "") -> bytes:
    df = df.copy() if df is not None else pd.DataFrame()
    k = turnover_kpis(df)
    insights = build_turnover_insights(df)
    now = datetime.now()
    period = k.get("period") or period_label
    params = f"Штуки{' и рубли по закупке' if k.get('has_money') else ''} · дата отчёта: {now:%d.%m.%Y} · период продаж: {period}"
    if scope:
        params += f" · {scope}"

    wb = Workbook()
    toc = wb.active
    toc.title = TOC_SHEET_NAME

    sheets_desc = [
        ("Выводы", "Выводы и рекомендации по оборачиваемости"),
        ("По статусам", "Распределение запаса по скорости оборачиваемости"),
        ("По категориям", "Оборачиваемость и «замороженный» запас по категориям"),
        ("По складам", "Где лежит запас: остаток, неликвид и медленный товар по складам"),
        ("ABC × оборачиваемость", "Доступный остаток по классам ABC и скорости оборачиваемости"),
        ("Неликвид", "Позиции без продаж и с оборачиваемостью больше года — к разбору"),
        ("Оборачиваемость SKU", "Все позиции матрицы: остаток, продажи, приходы, оборачиваемость"),
    ]

    # ---------------------------------------------------------- Оглавление
    _setup(toc, landscape=False)
    width = 8
    _write(toc, 1, 1, "АНАЛИЗ ОБОРАЧИВАЕМОСТИ ТОВАРОВ", font=_font(14, True))
    toc.row_dimensions[1].height = 24
    _write(toc, 2, 1, "Ассортиментная матрица · остатки, продажи и приходы", font=_font(9, color=TEXT_2))
    _write(toc, 3, 1, params, font=_font(9, color=MUTED))
    for cc in range(1, width + 1):
        toc.cell(row=4, column=cc).border = Border(bottom=_side(NAVY, "medium"))
    toc.row_dimensions[4].height = 6

    days = k.get("company_turnover_days")
    cards = [
        ("ОБОРАЧИВАЕМОСТЬ, ДН.", days, FMT_QTY, "весь доступный запас"),
        ("ДОСТУПНЫЙ ОСТАТОК, ШТ.", k.get("stock_units"), FMT_QTY, f"{k.get('sku_with_stock', 0)} SKU"),
        ("НЕЛИКВИД, ШТ.", k.get("dead_units"), FMT_QTY, f"{k.get('dead_sku', 0)} SKU · {k.get('dead_share', 0):.0%} запаса"),
        (
            "ОСТАТОК ПО ЗАКУПКЕ, ₽" if k.get("has_money") else "ЗАМОРОЖЕНО ЗАПАСА",
            k.get("stock_value") if k.get("has_money") else (k.get("frozen_share") or 0) * 100,
            FMT_MONEY if k.get("has_money") else FMT_PCT,
            "по цене последнего прихода" if k.get("has_money") else "неликвид + > 365 дн.",
        ),
    ]
    col = 1
    for i, (label, value, fmt, note) in enumerate(cards):
        c1, c2 = col, col + 1
        for rr in (6, 7, 8):
            toc.merge_cells(start_row=rr, start_column=c1, end_row=rr, end_column=c2)
        fill = SURFACE_4 if i == 0 else None
        _write(toc, 6, c1, label, font=_font(8, True, NAVY_3 if i == 0 else MUTED), fill=fill)
        _write(toc, 7, c1, value, font=_font(16, True, NAVY_3), fill=fill, fmt=fmt)
        _write(toc, 8, c1, note, font=_font(8, color=MUTED), fill=fill)
        for cc in (c1, c2):
            if fill:
                for rr in (6, 7, 8):
                    toc.cell(row=rr, column=cc).fill = _fill(fill)
            toc.cell(row=6, column=cc).border = Border(top=_side(NAVY, "medium"),
                                                       left=_side() if cc == c1 else None,
                                                       right=_side() if cc == c2 else None)
            toc.cell(row=7, column=cc).border = Border(left=_side() if cc == c1 else None,
                                                       right=_side() if cc == c2 else None)
            toc.cell(row=8, column=cc).border = Border(bottom=_side(),
                                                       left=_side() if cc == c1 else None,
                                                       right=_side() if cc == c2 else None)
        col += 2
    toc.row_dimensions[6].height = 16
    toc.row_dimensions[7].height = 26
    toc.row_dimensions[8].height = 16

    _write(toc, 10, 1, "СОДЕРЖАНИЕ ОТЧЁТА", font=_font(12, True))
    _write(toc, 11, 1, "Щёлкните на названии листа, чтобы перейти. На каждом листе есть кнопка возврата в оглавление.",
           font=_font(8, color=MUTED, italic=True))
    r = 12
    for name, desc in sheets_desc:
        toc.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        toc.merge_cells(start_row=r, start_column=4, end_row=r, end_column=width)
        c = _write(toc, r, 1, "›  " + name, fill=SURFACE_4, indent=1)
        c.hyperlink = f"#'{name}'!A1"
        c.font = _font(10, True, NAVY_3)
        for cc in (2, 3):
            toc.cell(row=r, column=cc).fill = _fill(SURFACE_4)
        _write(toc, r, 4, desc, font=_font(9, color=TEXT_2), wrap=True)
        if r % 2 == 0:
            for cc in range(4, width + 1):
                toc.cell(row=r, column=cc).fill = _fill(ZEBRA_ROW)
        for cc in range(1, width + 1):
            cell = toc.cell(row=r, column=cc)
            cell.border = Border(bottom=_side(), right=_side() if cc == 3 else None)
        toc.row_dimensions[r].height = 28
        r += 1
    _write(toc, r + 1, 1, f"Файл сформирован автоматически · {now:%d.%m.%Y %H:%M}", font=_font(8, color=MUTED))
    for j in range(1, width + 1):
        toc.column_dimensions[get_column_letter(j)].width = 13

    # -------------------------------------------------------------- Выводы
    ws = wb.create_sheet("Выводы")
    _setup(ws, landscape=False)
    r = _sheet_header(ws, 3, "ВЫВОДЫ И РЕКОМЕНДАЦИИ", "Что показывает оборачиваемость и что с этим делать", params)
    level_style = {
        "bad": ("Требует действия", EXPENSE, OCCUPIED_BG),
        "warn": ("Внимание", WARN, WARN_BG),
        "ok": ("Норма", NAVY_3, SURFACE_4),
        "info": ("Информация", INFO, INFO_BG),
    }
    rows = []
    for level, text in insights["findings"]:
        rows.append([level_style[level][0], text])
    end = _table(ws, r, [("Оценка", None, "left"), ("Вывод", None, "left")], rows,
                 widths=[30, 110], freeze=False)
    for i, (level, _) in enumerate(insights["findings"]):
        cell = ws.cell(row=r + 1 + i, column=1)
        cell.font = _font(10, True, level_style[level][1])
        cell.fill = _fill(level_style[level][2])
        ws.cell(row=r + 1 + i, column=2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r + 1 + i].height = 32

    r = end + 2
    _write(ws, r, 1, "РЕКОМЕНДАЦИИ", font=_font(12, True))
    r += 1
    end = _table(ws, r, [("Направление", None, "left"), ("Что сделать", None, "left")],
                 [[t, x] for t, x in insights["actions"]], freeze=False)
    for i in range(len(insights["actions"])):
        ws.cell(row=r + 1 + i, column=1).font = _font(10, True, NAVY_3)
        ws.cell(row=r + 1 + i, column=1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row=r + 1 + i, column=2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r + 1 + i].height = 32

    r = end + 2
    notes = [
        "Оборачиваемость, дн. = доступный остаток / среднедневные продажи. Для товаров, появившихся внутри периода, дни считаются с первой продажи или прихода (не меньше 30).",
        "Реализация партии = продано с последнего прихода / (продано с прихода + текущий остаток).",
        "Неликвид — есть остаток, но ни одной продажи за период и последний приход старше 60 дней.",
    ]
    for t in notes:
        _write(ws, r, 1, t, font=_font(8, color=MUTED))
        r += 1

    # ---------------------------------------------------------- По статусам
    ws = wb.create_sheet("По статусам")
    _setup(ws)
    st = turnover_by_status(df) if k.get("has_data") else pd.DataFrame()
    cols = [("Статус оборачиваемости", None, "left"), ("SKU", FMT_QTY, "right"),
            ("Остаток, шт.", FMT_QTY, "right"), ("Доля остатка", FMT_PCT, "right"),
            ("Остаток по закупке, ₽", FMT_MONEY, "right")]
    r = _sheet_header(ws, len(cols), "ЗАПАС ПО СКОРОСТИ ОБОРАЧИВАЕМОСТИ", "Сколько SKU и штук в каждом статусе", params)
    rows = [[x.status, x.sku, x.stock, x.stock_share * 100, _n(x.value)] for x in st.itertuples()] if not st.empty else []
    total = ["ИТОГО", st["sku"].sum() if not st.empty else None, st["stock"].sum() if not st.empty else None,
             100.0 if not st.empty else None, _n(st["value"].sum(min_count=1)) if not st.empty else None]
    _table(ws, r, cols, rows, total=total, widths=[38, 12, 16, 16, 22])

    # -------------------------------------------------------- По категориям
    ws = wb.create_sheet("По категориям")
    _setup(ws)
    cat = turnover_by_category(df) if k.get("has_data") else pd.DataFrame()
    cols = [("Категория", None, "left"), ("SKU с остатком", FMT_QTY, "right"), ("Остаток, шт.", FMT_QTY, "right"),
            ("Продано за период, шт.", FMT_QTY, "right"),
            ("Продажи в день, шт.", FMT_DEC, "right"), ("Оборачиваемость, дн.", FMT_QTY, "right"),
            ("Неликвид, шт.", FMT_QTY, "right"), ("Заморожено, %", FMT_PCT, "right"),
            ("Остаток по закупке, ₽", FMT_MONEY, "right")]
    r = _sheet_header(ws, len(cols), "ОБОРАЧИВАЕМОСТЬ ПО КАТЕГОРИЯМ",
                      "Заморожено = неликвид + позиции с оборачиваемостью больше года", params)
    rows = [[x.cat_name, x.sku, x.stock, x.sold, x.daily, _n(x.turnover_days), x.dead, x.frozen_share * 100, _n(x.value)]
            for x in cat.itertuples()] if not cat.empty else []
    if not cat.empty:
        tot_stock, tot_daily = cat["stock"].sum(), cat["daily"].sum()
        total = ["ИТОГО", cat["sku"].sum(), tot_stock, cat["sold"].sum(), tot_daily,
                 (tot_stock / tot_daily) if tot_daily else None, cat["dead"].sum(),
                 (cat["frozen"].sum() / tot_stock * 100) if tot_stock else None, _n(cat["value"].sum(min_count=1))]
    else:
        total = None
    _table(ws, r, cols, rows, total=total, widths=[34, 14, 14, 16, 16, 16, 14, 14, 20], autofilter=True)

    # -------------------------------------------------------- По складам
    ws = wb.create_sheet("По складам")
    _setup(ws)
    wh = turnover_by_warehouse(df) if k.get("has_data") else pd.DataFrame()
    cols = [("Склад", None, "left"), ("SKU", FMT_QTY, "right"), ("Остаток, шт.", FMT_QTY, "right"),
            ("Доля запаса, %", FMT_PCT, "right"), ("Неликвид, шт.", FMT_QTY, "right"),
            ("Очень медленные, шт.", FMT_QTY, "right"), ("Заморожено, %", FMT_PCT, "right")]
    r = _sheet_header(ws, len(cols), "ЗАПАС ПО СКЛАДАМ",
                      "Статус оборачиваемости — по товару в целом; высокая доля в салонах часто означает выставочные образцы",
                      params)
    rows = [[x.warehouse, x.sku, x.stock, x.share * 100, x.dead, x.very_slow, x.frozen_share * 100]
            for x in wh.itertuples()] if not wh.empty else []
    total = None
    if not wh.empty:
        t_stock = wh["stock"].sum()
        total = ["ИТОГО", None, t_stock, 100.0, wh["dead"].sum(), wh["very_slow"].sum(),
                 (wh["frozen"].sum() / t_stock * 100) if t_stock else None]
    _table(ws, r, cols, rows, total=total, widths=[36, 10, 14, 14, 14, 18, 14], autofilter=True)

    # ------------------------------------------------ ABC × оборачиваемость
    ws = wb.create_sheet("ABC × оборачиваемость")
    _setup(ws)
    m = abc_turnover_matrix(df)
    cols = [("ABC", None, "left")] + [(c, FMT_QTY, "right") for c in m["cols"]] + [("ИТОГО, шт.", FMT_QTY, "right")]
    r = _sheet_header(ws, len(cols), "ABC × ОБОРАЧИВАЕМОСТЬ",
                      "Доступный остаток, шт.; A/B с медленной оборачиваемостью — лишний запас в ходовом товаре",
                      params)
    rows = [[a] + list(vals) + [sum(vals)] for a, vals in zip(m["rows"], m["stock"])]
    total = None
    if rows:
        total = ["ИТОГО"] + [sum(col) for col in zip(*m["stock"])] + [sum(map(sum, m["stock"]))]
    end = _table(ws, r, cols, rows, total=total, widths=[14] + [18] * (len(cols) - 1))
    r = end + 2
    _write(ws, r, 1, "Количество SKU", font=_font(12, True))
    cols_sku = [("ABC", None, "left")] + [(c, FMT_QTY, "right") for c in m["cols"]]
    _table(ws, r + 1, cols_sku, [[a] + list(v) for a, v in zip(m["rows"], m["sku"])], freeze=False)

    # ------------------------------------------------- SKU (общий и неликвид)
    sku_cols = [
        ("fullname", "Номенклатура", None, "left", 45),
        ("article", "Артикул", None, "left", 16),
        ("manu", "Производитель", None, "left", 20),
        ("cat_name", "Категория", None, "left", 22),
        ("abc", "ABC", None, "center", 8),
        ("turnover_status", "Статус оборачиваемости", None, "left", 28),
        ("stock_available", "Доступно, шт.", FMT_QTY, "right", 12),
        ("stock_ordered", "Заказано, шт.", FMT_QTY, "right", 12),
        ("quant", "Продано за период, шт.", FMT_QTY, "right", 14),
        ("avg_daily_sales", "Продажи в день", FMT_DEC, "right", 12),
        ("turnover_days", "Оборачиваемость, дн.", FMT_QTY, "right", 14),
        ("active_days", "Дней в продаже", FMT_QTY, "right", 12),
        ("turns_per_year", "Оборотов в год", FMT_DEC, "right", 12),
        ("first_receipt_date", "Первый приход", FMT_DATE, "center", 12),
        ("last_receipt_date", "Последний приход", FMT_DATE, "center", 12),
        ("days_since_receipt", "Дней с прихода", FMT_QTY, "right", 12),
        ("last_receipt_qty", "Посл. партия, шт.", FMT_QTY, "right", 12),
        ("sold_since_receipt", "Продано с прихода, шт.", FMT_QTY, "right", 14),
        ("sell_through", "Реализация партии, %", FMT_PCT, "right", 13),
        ("receipt_qty_period", "Пришло за период, шт.", FMT_QTY, "right", 13),
        ("purchase_price", "Цена закупки, ₽", FMT_MONEY, "right", 13),
        ("stock_value_purchase", "Остаток по закупке, ₽", FMT_MONEY, "right", 16),
        ("stock_status", "Статус запаса (ROP)", None, "left", 22),
        ("order_need", "Нужно заказать, шт.", FMT_QTY, "right", 13),
    ]
    sku_cols = [c for c in sku_cols if c[0] in df.columns]

    def _sku_rows(frame):
        out = []
        for rec in frame.to_dict("records"):
            row = []
            for field, _, fmt, _, _ in sku_cols:
                v = rec.get(field)
                if fmt == FMT_DATE:
                    v = _to_date(v)
                elif fmt == FMT_PCT:
                    v = None if _n(v) is None else _n(v) * 100
                elif fmt is not None:
                    v = _n(v)
                elif isinstance(v, float) and np.isnan(v):
                    v = None
                row.append(v)
            out.append(row)
        return out

    for name, title, sub, frame in (
        (
            "Неликвид",
            "НЕЛИКВИД И ОЧЕНЬ МЕДЛЕННЫЕ ПОЗИЦИИ",
            "Отсортировано по остатку: сверху — то, что держит больше всего запаса",
            df[df.get("turnover_status", pd.Series(dtype=str)).isin([STATUS_DEAD, STATUS_VERY_SLOW])]
            .sort_values("stock_available", ascending=False) if k.get("has_data") else df,
        ),
        (
            "Оборачиваемость SKU",
            "ОБОРАЧИВАЕМОСТЬ ПО ПОЗИЦИЯМ",
            "Все позиции матрицы с остатком или продажами за период",
            df.sort_values("stock_available", ascending=False) if "stock_available" in df.columns else df,
        ),
    ):
        ws = wb.create_sheet(name)
        _setup(ws)
        cols = [(t, f, a) for _, t, f, a, _ in sku_cols]
        r = _sheet_header(ws, len(cols), title, sub, params)
        rows = _sku_rows(frame)
        total = None
        if rows:
            total = []
            for field, t, f, a, _ in sku_cols:
                if field in ("stock_available", "stock_ordered", "quant", "stock_value_purchase",
                             "last_receipt_qty", "sold_since_receipt", "receipt_qty_period", "order_need"):
                    total.append(_n(pd.to_numeric(frame[field], errors="coerce").sum(min_count=1)))
                elif field == "fullname":
                    total.append(f"ИТОГО ({len(rows)} SKU)")
                else:
                    total.append(None)
        _table(ws, r, cols, rows, total=total, widths=[w for *_, w in sku_cols], autofilter=True)

    order = [TOC_SHEET_NAME] + [n for n, _ in sheets_desc]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]
    for sheet in wb.worksheets:
        for view in sheet.views.sheetView:
            view.tabSelected = False
    wb.active = 0

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
