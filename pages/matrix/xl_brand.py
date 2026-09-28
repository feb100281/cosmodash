# pages/matrix/xl_brand.py
"""Общий стиль Excel-выгрузок матрицы."""
from __future__ import annotations

from datetime import datetime
from typing import Iterable, Optional, Sequence

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

NAVY = "2F6656"
NAVY_2 = "3D7A67"
NAVY_3 = "1F5E4E"
TEXT = "1F1F1F"
TEXT_2 = "4A4A4A"
MUTED = "8A8A8A"
SURFACE_2 = "F3F8F6"
SURFACE_3 = "EDF5F1"
SURFACE_4 = "E7F1ED"
TOTAL_ROW = "E7F1ED"
ZEBRA_ROW = "F7F7F7"
LINE = "D9D9D9"
EXPENSE = "7B4437"
INFO = "2F75B5"
INFO_BG = "EAF2FB"
WARN = "9A6100"
WARN_BG = "FDF3DE"
OCCUPIED_BG = "F6E9E4"

FONT = "Roboto"

FMT_MONEY = '#,##0;(#,##0);"–"'
FMT_QTY = '#,##0;(#,##0);"–"'
FMT_DEC = '#,##0.00;(#,##0.00);"–"'
FMT_SHARE = '0.0%;(0.0%);"–"'
FMT_DATE = "dd.mm.yyyy"

TOC_SHEET_NAME = "Оглавление"


def fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def side(color: str = LINE, style: str = "thin") -> Side:
    return Side(style=style, color=color)


def font(size: int = 10, bold: bool = False, color: str = TEXT, italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=None)


def page_setup(ws, landscape: bool = True) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True


def toc_button(ws, row: int, width: int, toc_name: str = TOC_SHEET_NAME) -> None:
    width = max(width, 2)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    c = ws.cell(row=row, column=1, value="←  Оглавление")
    c.hyperlink = f"#'{toc_name}'!A1"
    c.font = Font(name=FONT, size=9, bold=True, color=NAVY_3, underline=None)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for cc in range(1, width + 1):
        cell = ws.cell(row=row, column=cc)
        cell.fill = fill(SURFACE_4)
        cell.border = Border(bottom=side())
    ws.row_dimensions[row].height = 16


def style_table_sheet(
    ws,
    *,
    title: str,
    subtitle: str,
    toc_name: Optional[str] = TOC_SHEET_NAME,
    header_row: int = 1,
    freeze_col: int = 2,
    key_headers: Iterable[str] = (),
    wrap_headers: Iterable[str] = (),
) -> int:
    """
    Оформляет лист, куда уже записан DataFrame (шапка в header_row).
    Сверху добавляются 3 строки: оглавление, заголовок, подзаголовок.
    Возвращает новый номер строки шапки.
    """
    ws.insert_rows(1, 3)
    header_row += 3
    max_col = max(ws.max_column, 2)
    max_row = ws.max_row

    page_setup(ws)
    if toc_name:
        toc_button(ws, 1, max_col, toc_name)
    ws.cell(row=2, column=1, value=title).font = font(14, True)
    ws.row_dimensions[2].height = 24
    c = ws.cell(row=3, column=1, value=subtitle)
    c.font = font(9, color=MUTED)
    for cc in range(1, max_col + 1):
        ws.cell(row=3, column=cc).border = Border(bottom=side(NAVY, "medium"))
    ws.row_dimensions[3].height = 18

    keys = {str(h).strip().lower() for h in key_headers}
    wraps = {str(h).strip().lower() for h in wrap_headers}

    for cc in range(1, max_col + 1):
        cell = ws.cell(row=header_row, column=cc)
        is_key = str(cell.value or "").strip().lower() in keys
        cell.fill = fill(NAVY_3 if is_key else NAVY)
        cell.font = font(10, True, "FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=side(), right=side(), top=side(), bottom=side())
    ws.row_dimensions[header_row].height = 32

    headers = {cc: str(ws.cell(row=header_row, column=cc).value or "").strip().lower()
               for cc in range(1, max_col + 1)}

    for r in range(header_row + 1, max_row + 1):
        for cc in range(1, max_col + 1):
            cell = ws.cell(row=r, column=cc)
            cell.font = font(10)
            cell.border = Border(bottom=side(), right=side() if cc < max_col else None)
            num = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            cell.alignment = Alignment(
                horizontal="right" if num else "left",
                vertical="center",
                wrap_text=headers[cc] in wraps,
            )
            if headers[cc] in keys:
                cell.fill = fill(SURFACE_3)

    apply_zebra(ws, header_row + 1, max_row, 1, max_col)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=freeze_col)
    if max_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{max_row}"
    return header_row


def apply_zebra(ws, first_row: int, last_row: int, first_col: int, last_col: int) -> None:
    for r in range(first_row, last_row + 1):
        if r % 2:
            continue
        for cc in range(first_col, last_col + 1):
            cell = ws.cell(row=r, column=cc)
            if not cell.fill.fill_type:
                cell.fill = fill(ZEBRA_ROW)


def set_number_format(ws, header_row: int, headers: Sequence[str], fmt: str,
                      *, prefix: bool = False) -> None:
    targets = {h.strip().lower() for h in headers}
    for cc in range(1, ws.max_column + 1):
        h = str(ws.cell(row=header_row, column=cc).value or "").strip().lower()
        hit = any(h.startswith(t) for t in targets) if prefix else h in targets
        if not hit:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row=r, column=cc)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="right", vertical="center")


def build_toc(
    wb,
    *,
    title: str,
    subtitle: str,
    params: str,
    sheets: Sequence[tuple[str, str]],
    cards: Sequence[tuple[str, object, str, str]] = (),
    children: dict[str, Sequence[str]] | None = None,
    toc_name: str = TOC_SHEET_NAME,
):
    """
    sheets: [(имя листа, описание)]
    cards:  [(подпись, значение, формат, примечание)] — до 4 шт.
    children: {родительский лист: [дочерние листы]} — выводятся списком под родителем.
    """
    if toc_name in wb.sheetnames:
        wb.remove(wb[toc_name])
    ws = wb.create_sheet(toc_name, 0)
    page_setup(ws, landscape=False)
    width = 8

    ws.cell(row=1, column=1, value=title.upper()).font = font(14, True)
    ws.row_dimensions[1].height = 24
    ws.cell(row=2, column=1, value=subtitle).font = font(9, color=TEXT_2)
    ws.cell(row=3, column=1, value=params).font = font(9, color=MUTED)
    for cc in range(1, width + 1):
        ws.cell(row=4, column=cc).border = Border(bottom=side(NAVY, "medium"))
    ws.row_dimensions[4].height = 6

    r = 6
    if cards:
        col = 1
        for i, (label, value, fmt, note) in enumerate(cards[:4]):
            c1, c2 = col, col + 1
            for rr in (r, r + 1, r + 2):
                ws.merge_cells(start_row=rr, start_column=c1, end_row=rr, end_column=c2)
            bg = SURFACE_4 if i == 0 else None
            for rr, (v, f, fm) in zip(
                (r, r + 1, r + 2),
                ((label, font(8, True, NAVY_3 if i == 0 else MUTED), None),
                 (value, font(16, True, NAVY_3), fmt),
                 (note, font(8, color=MUTED), None)),
            ):
                cell = ws.cell(row=rr, column=c1, value=v)
                cell.font = f
                if fm:
                    cell.number_format = fm
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for cc in (c1, c2):
                for rr in (r, r + 1, r + 2):
                    cell = ws.cell(row=rr, column=cc)
                    if bg:
                        cell.fill = fill(bg)
                    cell.border = Border(
                        top=side(NAVY, "medium") if rr == r else None,
                        bottom=side() if rr == r + 2 else None,
                        left=side() if cc == c1 else None,
                        right=side() if cc == c2 else None,
                    )
            col += 2
        ws.row_dimensions[r].height = 16
        ws.row_dimensions[r + 1].height = 26
        ws.row_dimensions[r + 2].height = 16
        r += 4

    ws.cell(row=r, column=1, value="СОДЕРЖАНИЕ ОТЧЁТА").font = font(12, True)
    ws.cell(row=r + 1, column=1,
            value="Щёлкните на названии листа, чтобы перейти. На каждом листе есть кнопка возврата в оглавление."
            ).font = font(8, color=MUTED, italic=True)
    r += 2

    def _row(name, desc, child=False):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=width)
        c = ws.cell(row=r, column=1, value=("      " if child else "›  ") + name)
        c.hyperlink = f"#'{name}'!A1"
        c.font = font(9 if child else 10, not child, NAVY_3)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        d = ws.cell(row=r, column=4, value=desc)
        d.font = font(9, color=TEXT_2)
        d.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for cc in range(1, width + 1):
            cell = ws.cell(row=r, column=cc)
            if cc <= 3:
                cell.fill = fill(SURFACE_2 if child else SURFACE_4)
            elif r % 2 == 0:
                cell.fill = fill(ZEBRA_ROW)
            cell.border = Border(bottom=side(), right=side() if cc == 3 else None)
        ws.row_dimensions[r].height = 20 if child else 28
        r += 1

    for name, desc in sheets:
        if name not in wb.sheetnames:
            continue
        _row(name, desc)
        for child in (children or {}).get(name, []):
            if child in wb.sheetnames:
                _row(child, "", child=True)

    ws.cell(row=r + 1, column=1,
            value=f"Файл сформирован автоматически · {datetime.now():%d.%m.%Y %H:%M}").font = font(8, color=MUTED)
    for j in range(1, width + 1):
        ws.column_dimensions[get_column_letter(j)].width = 13
    return ws


def finalize(wb, order: Sequence[str] = ()) -> None:
    if order:
        ordered = [wb[n] for n in order if n in wb.sheetnames]
        ordered += [s for s in wb.worksheets if s not in ordered]
        wb._sheets = ordered
    for sheet in wb.worksheets:
        for view in sheet.views.sheetView:
            view.tabSelected = False
    wb.active = 0
