# pages/cats/cats_export.py
"""Выгрузки раздела «Категории»: Excel-отчёт и справка (PDF)."""
from __future__ import annotations

import html
from datetime import datetime
from io import BytesIO

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter

from pages.matrix.xl_brand import (
    EXPENSE,
    FMT_MONEY,
    FMT_QTY,
    FMT_SHARE,
    MUTED,
    NAVY,
    NAVY_3,
    TEXT_2,
    TOTAL_ROW,
    build_toc,
    fill,
    finalize,
    font,
    page_setup,
    side,
    style_table_sheet,
)

from .analysis import VALS_DICT, build_narrative, level_table

LEVEL_SHEETS = [
    ("parent", "Группы", ["Группа"]),
    ("cat", "Категории", ["Группа", "Категория"]),
    ("subcat", "Подкатегории", ["Группа", "Категория", "Подкатегория"]),
]


def _columns(data):
    cur, ref = data["cur_label"], data["ref_label"]
    return [
        ("amount_cur", f"Выручка, {cur}", FMT_MONEY),
        ("amount_ref", f"Выручка, {ref}", FMT_MONEY),
        ("amount_var", "Выручка, Δ", FMT_MONEY),
        ("amount_pct", "Выручка, Δ%", FMT_SHARE),
        ("share", "Доля выручки", FMT_SHARE),
        ("quant_cur", f"Кол-во, {cur}", FMT_QTY),
        ("quant_ref", f"Кол-во, {ref}", FMT_QTY),
        ("quant_var", "Кол-во, Δ", FMT_QTY),
        ("cr_cur", f"Возвраты, {cur}", FMT_MONEY),
        ("cr_ref", f"Возвраты, {ref}", FMT_MONEY),
        ("ret_share_cur", "Доля возвратов", FMT_SHARE),
    ]


def _write_level(wb, data, level, sheet, key_titles, params):
    t = level_table(data["df"], level)
    cols = _columns(data)
    ws = wb.create_sheet(sheet)
    keys = [k for k in ("parent", "cat", "subcat") if k in t.columns][: len(key_titles)]

    ws.append(key_titles + [title for _, title, _ in cols])
    for rec in t.to_dict("records"):
        row = [rec[k] for k in keys]
        for field, _, _ in cols:
            v = rec.get(field)
            row.append(None if v is None or (isinstance(v, float) and np.isnan(v)) else float(v))
        ws.append(row)

    total = ["ИТОГО"] + [""] * (len(keys) - 1)
    sums = {f: float(t[f].sum()) for f, _, _ in cols if not f.endswith(("_pct", "share", "ret_share_cur"))}
    for field, _, _ in cols:
        if field in sums:
            total.append(sums[field])
        elif field == "amount_pct":
            total.append(sums["amount_var"] / sums["amount_ref"] if sums["amount_ref"] else None)
        elif field == "share":
            total.append(1.0)
        elif field == "ret_share_cur":
            dt = float(t["dt_cur"].sum())
            total.append(sums["cr_cur"] / dt if dt else None)
        else:
            total.append(None)

    hr = style_table_sheet(
        ws,
        title=f"{sheet.upper()}: {data['cur_label']} против {data['ref_gen']}",
        subtitle=params,
        header_row=1,
        freeze_col=len(keys) + 1,
        key_headers=[f"Выручка, {data['cur_label']}", "Выручка, Δ"],
    )
    for j, (_, _, fmt) in enumerate(cols, start=len(keys) + 1):
        for r in range(hr + 1, ws.max_row + 1):
            ws.cell(row=r, column=j).number_format = fmt
            ws.cell(row=r, column=j).alignment = Alignment(horizontal="right", vertical="center")

    var_cols = [j for j, (f, _, _) in enumerate(cols, start=len(keys) + 1) if f in ("amount_var", "amount_pct", "quant_var")]
    for r in range(hr + 1, ws.max_row + 1):
        for j in var_cols:
            cell = ws.cell(row=r, column=j)
            if isinstance(cell.value, (int, float)) and cell.value < 0:
                cell.font = font(10, color=EXPENSE)

    r = ws.max_row + 1
    for j, v in enumerate(total, start=1):
        cell = ws.cell(row=r, column=j, value=v)
        cell.font = font(10, True)
        cell.fill = fill(TOTAL_ROW)
        cell.border = Border(top=side(NAVY, "medium"), bottom=side(NAVY, "double"))
        if j > len(keys):
            cell.number_format = cols[j - len(keys) - 1][2]
            cell.alignment = Alignment(horizontal="right", vertical="center")
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(ws.max_column)}{r - 1}"

    widths = [22, 30, 34][: len(keys)] + [15] * len(cols)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _write_reference(wb, sections, title, params):
    ws = wb.create_sheet("Справка")
    page_setup(ws, landscape=False)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    from pages.matrix.xl_brand import toc_button
    toc_button(ws, 1, 2)
    ws.cell(row=2, column=1, value=title.upper()).font = font(14, True)
    ws.cell(row=3, column=1, value=params).font = font(9, color=MUTED)
    for cc in (1, 2):
        ws.cell(row=3, column=cc).border = Border(bottom=side(NAVY, "medium"))
    r = 5
    for sec in sections:
        c = ws.cell(row=r, column=1, value=sec["title"])
        c.font = font(11, True, NAVY_3)
        r += 1
        for item in sec["items"]:
            b = ws.cell(row=r, column=1, value="•")
            b.font = font(10, color=NAVY)
            b.alignment = Alignment(vertical="top", horizontal="center")
            c = ws.cell(row=r, column=2, value=item)
            c.font = font(10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 15 * (len(item) // 105 + 1))
            r += 1
        r += 1
    ws.cell(row=r, column=1, value=f"Сформировано автоматически · {datetime.now():%d.%m.%Y %H:%M}").font = font(8, color=MUTED)


def build_cats_excel_bytes(data: dict, val: str = "amount", scope: str = "") -> bytes:
    sections = build_narrative(data, val)
    groups = level_table(data["df"], "parent") if not data["df"].empty else None
    params = (f"Рубли и штуки · {data['cur_label']} против {data['ref_gen']} · "
              f"выгрузка {datetime.now():%d.%m.%Y}" + (f" · {scope}" if scope else ""))

    wb = Workbook()
    wb.remove(wb.active)
    _write_reference(wb, sections, "Справка по категориям", params)
    for level, sheet, keys in LEVEL_SHEETS:
        _write_level(wb, data, level, sheet, keys, params)

    cards = []
    if groups is not None:
        cur, ref = groups["amount_cur"].sum(), groups["amount_ref"].sum()
        var = cur - ref
        cards = [
            (f"ВЫРУЧКА, {data['cur_label'].upper()}", cur, FMT_MONEY, "текущий месяц"),
            (f"ВЫРУЧКА, {data['ref_label'].upper()}", ref, FMT_MONEY, "база сравнения"),
            ("ИЗМЕНЕНИЕ, ₽", var, FMT_MONEY, f"{(var / ref if ref else 0):+.1%}".replace(".", ",")),
            ("КОЛИЧЕСТВО, ШТ.", groups["quant_cur"].sum(), FMT_QTY, data["cur_label"]),
        ]
    build_toc(
        wb,
        title="Анализ категорий",
        subtitle=f"Сравнение месяцев: {data['cur_label']} против {data['ref_gen']}",
        params=params,
        cards=cards,
        sheets=[
            ("Справка", "Выводы обычным языком: что выросло, что просело и что с этим делать"),
            ("Группы", "Выручка, количество и возвраты по группам"),
            ("Категории", "То же по категориям внутри групп"),
            ("Подкатегории", "То же по подкатегориям"),
        ],
    )
    finalize(wb, ["Оглавление", "Справка", "Группы", "Категории", "Подкатегории"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Справка (PDF)
# ---------------------------------------------------------------------------

_CSS = """
@page { size: A4; margin: 18mm 16mm 16mm 16mm;
        @bottom-right { content: counter(page); font-size: 8pt; color: #8A8A8A; } }
body { font-family: Roboto, "Helvetica Neue", Arial, sans-serif; color: #1F1F1F; font-size: 10pt; line-height: 1.45; }
.kicker { font-size: 8pt; letter-spacing: 1.5px; text-transform: uppercase; color: #8A8A8A; }
h1 { font-size: 20pt; color: #1F5E4E; margin: 4px 0 6px; }
.rule { height: 2px; background: #2F6656; width: 60px; margin: 8px 0 14px; }
.params { color: #4A4A4A; font-size: 9pt; margin-bottom: 14px; }
.cards { display: flex; gap: 8px; margin: 10px 0 18px; }
.card { flex: 1; border: 1px solid #D9D9D9; border-top: 2px solid #2F6656; padding: 8px 10px; }
.card.main { background: #E7F1ED; }
.card .l { font-size: 7.5pt; color: #8A8A8A; text-transform: uppercase; font-weight: 700; }
.card .v { font-size: 15pt; font-weight: 700; color: #1F5E4E; margin-top: 2px; }
.card .n { font-size: 8pt; color: #8A8A8A; }
h2 { font-size: 12pt; color: #1F5E4E; border-bottom: 1px solid #D9D9D9; padding-bottom: 3px; margin: 16px 0 6px; }
ul { margin: 0; padding-left: 16px; }
li { margin: 3px 0; }
table { width: 100%; border-collapse: collapse; margin-top: 8px; font-size: 8.5pt; }
th { background: #2F6656; color: #fff; text-align: left; padding: 4px 6px; font-weight: 700; }
td { padding: 3px 6px; border-bottom: 1px solid #E6E6E6; }
td.n, th.n { text-align: right; }
tr:nth-child(even) td { background: #F7F7F7; }
.neg { color: #7B4437; }
.foot { margin-top: 18px; font-size: 8pt; color: #8A8A8A; }
"""


def _money(v):
    from .analysis import fmt_value
    return fmt_value(v, "amount")


def build_cats_reference_html(data: dict, val: str = "amount", scope: str = "") -> str:
    sections = build_narrative(data, val)
    e = html.escape
    parts = [
        f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>",
        "<div class='kicker'>Справка · анализ категорий</div>",
        f"<h1>{e(data['cur_label'].capitalize())} против {e(data['ref_gen'])}</h1><div class='rule'></div>",
        f"<div class='params'>Метрика выводов: {e(VALS_DICT.get(val, val))}"
        + (f" · {e(scope)}" if scope else "") + "</div>",
    ]
    if not data["df"].empty:
        g = level_table(data["df"], "parent")
        cur, ref = g["amount_cur"].sum(), g["amount_ref"].sum()
        var = cur - ref
        pct = f"{(var / ref if ref else 0) * 100:+.1f}%".replace(".", ",").replace("-", "−")
        parts.append("<div class='cards'>")
        for cls, label, value, note in (
            ("main", f"Выручка, {data['cur_label']}", _money(cur), "текущий месяц"),
            ("", f"Выручка, {data['ref_label']}", _money(ref), "база сравнения"),
            ("", "Изменение", ("+" if var >= 0 else "−") + _money(abs(var)), pct),
            ("", "Количество", f"{g['quant_cur'].sum():,.0f} шт.".replace(",", " "), data["cur_label"]),
        ):
            parts.append(f"<div class='card {cls}'><div class='l'>{e(label)}</div>"
                         f"<div class='v'>{e(value)}</div><div class='n'>{e(note)}</div></div>")
        parts.append("</div>")

    for sec in sections:
        parts.append(f"<h2>{e(sec['title'])}</h2><ul>")
        parts += [f"<li>{e(i)}</li>" for i in sec["items"]]
        parts.append("</ul>")

    if not data["df"].empty:
        cats = level_table(data["df"], "cat")
        parts.append("<h2>Категории: цифры</h2><table><tr><th>Категория</th><th>Группа</th>"
                     f"<th class='n'>{e(data['cur_label'])}</th><th class='n'>{e(data['ref_label'])}</th>"
                     "<th class='n'>Δ</th><th class='n'>Δ%</th><th class='n'>Доля</th></tr>")
        from .analysis import fmt_pct
        for r in cats.itertuples():
            neg = " neg" if r.amount_var < 0 else ""
            parts.append(
                f"<tr><td>{e(r.cat)}</td><td>{e(r.parent)}</td>"
                f"<td class='n'>{e(_money(r.amount_cur))}</td><td class='n'>{e(_money(r.amount_ref))}</td>"
                f"<td class='n{neg}'>{e(('+' if r.amount_var >= 0 else '−') + _money(abs(r.amount_var)))}</td>"
                f"<td class='n{neg}'>{e(fmt_pct(r.amount_pct))}</td>"
                f"<td class='n'>{f'{r.share * 100:.1f}'.replace('.', ',')}%</td></tr>"
            )
        parts.append("</table>")

    parts.append(f"<div class='foot'>Сформировано автоматически · {datetime.now():%d.%m.%Y %H:%M}. "
                 "Сравниваются полные календарные месяцы: первый и последний месяц выбранного периода.</div>")
    parts.append("</body></html>")
    return "".join(parts)


def build_cats_reference_pdf(data: dict, val: str = "amount", scope: str = "") -> bytes:
    from weasyprint import HTML
    return HTML(string=build_cats_reference_html(data, val, scope)).write_pdf()
