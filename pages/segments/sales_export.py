# pages/segments/sales_export.py
"""Выгрузка всех продаж за период (построчно) с учётом выбора позиций."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

from data import ENGINE
import numpy as np

from pages.matrix.xl_brand import (
    FMT_DATE,
    FMT_SHARE,
    SURFACE_4 as _S4,
    TEXT_2,
    ZEBRA_ROW,
    FMT_MONEY,
    FMT_QTY,
    FONT,
    MUTED,
    NAVY,
    NAVY_3,
    SURFACE_4,
    TEXT,
    side,
)

COLUMNS = [
    ("date", "Дата", FMT_DATE, 12),
    ("fullname", "Номенклатура", None, 46),
    ("article", "Артикул", None, 16),
    ("barcode", "Штрихкод", None, 16),
    ("parent_cat", "Группа", None, 18),
    ("cat", "Категория", None, 22),
    ("subcat", "Подкатегория", None, 22),
    ("manu", "Производитель", None, 22),
    ("brend", "Бренд", None, 18),
    ("store_group", "Магазин (группа)", None, 22),
    ("store", "Подразделение", None, 26),
    ("agent", "Дизайнер", None, 22),
    ("manager", "Менеджер", None, 20),
    ("quant_dt", "Продано, шт.", FMT_QTY, 11),
    ("quant_cr", "Возвраты, шт.", FMT_QTY, 11),
    ("dt", "Продажи, ₽", FMT_MONEY, 14),
    ("cr", "Возвраты, ₽", FMT_MONEY, 14),
    ("amount", "Выручка, ₽", FMT_MONEY, 14),
]
FREEZE_AFTER = 2  # Дата + Номенклатура
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
          "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]


def fetch_sales(start: str, end: str, item_ids: list[int] | None = None, search: str | None = None) -> pd.DataFrame:
    where = ["s.date BETWEEN %(start)s AND %(end)s"]
    params = {"start": start, "end": end}
    if item_ids:
        where.append(f"s.item_id IN ({','.join(str(int(i)) for i in item_ids)})")
    if search:
        where.append("i.fullname LIKE %(q)s")
        params["q"] = f"%{search}%"
    q = f"""
        SELECT
            s.date,
            COALESCE(sg.name, 'Без магазина') AS store_group,
            COALESCE(st.name, '') AS store,
            COALESCE(parent.name, 'Без группы') AS parent_cat,
            COALESCE(cat.name, 'Без категории') AS cat,
            COALESCE(sc.name, 'Нет подкатегории') AS subcat,
            COALESCE(m.name, '') AS manu,
            COALESCE(b.name, '') AS brend,
            i.fullname,
            COALESCE(i.article, '') AS article,
            COALESCE(bc.barcode, '') AS barcode,
            COALESCE(a.report_name, 'Без дизайнера') AS agent,
            COALESCE(mg.report_name, '') AS manager,
            s.quant_dt, s.quant_cr, s.dt, s.cr, (s.dt - s.cr) AS amount
        FROM sales_salesdata AS s
        LEFT JOIN corporate_items AS i ON i.id = s.item_id
        LEFT JOIN corporate_itemmanufacturer AS m ON m.id = i.manufacturer_id
        LEFT JOIN corporate_itembrend AS b ON b.id = i.brend_id
        LEFT JOIN corporate_cattree AS cat ON cat.id = i.cat_id
        LEFT JOIN corporate_cattree AS parent ON parent.id = cat.parent_id
        LEFT JOIN corporate_subcategory AS sc ON sc.id = i.subcat_id
        LEFT JOIN corporate_barcode AS bc ON bc.id = s.barcode_id
        LEFT JOIN corporate_agents AS a ON a.id = s.agent_id
        LEFT JOIN corporate_managers AS mg ON mg.id = s.manager_id
        LEFT JOIN corporate_stores AS st ON st.id = s.store_id
        LEFT JOIN corporate_storegroups AS sg ON sg.id = st.gr_id
        WHERE {' AND '.join(where)}
        ORDER BY s.date, i.fullname
    """
    return pd.read_sql(q, ENGINE, params=params)


class _Styles:
    def __init__(self):
        self.title = Font(name=FONT, size=14, bold=True, color=TEXT)
        self.muted = Font(name=FONT, size=9, color=MUTED)
        self.body = Font(name=FONT, size=10, color=TEXT)
        self.bold = Font(name=FONT, size=10, bold=True, color=TEXT)
        self.link = Font(name=FONT, size=9, bold=True, color=NAVY_3, underline=None)
        self.link10 = Font(name=FONT, size=10, bold=True, color=NAVY_3, underline=None)
        self.hdr = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        self.card_label = Font(name=FONT, size=8, bold=True, color=MUTED)
        self.card_value = Font(name=FONT, size=16, bold=True, color=NAVY_3)
        self.hdr_fill = PatternFill("solid", start_color=NAVY, end_color=NAVY)
        self.zebra = PatternFill("solid", start_color=ZEBRA_ROW, end_color=ZEBRA_ROW)
        self.surface = PatternFill("solid", start_color=SURFACE_4, end_color=SURFACE_4)
        self.center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.right = Alignment(horizontal="right", vertical="center")
        self.left = Alignment(horizontal="left", vertical="center")
        self.body_border = Border(bottom=side(), right=side())
        self.hdr_border = Border(left=side(), right=side(), top=side(), bottom=side())
        self.brand_line = Border(bottom=side(NAVY, "medium"))
        self.total_border = Border(top=side(NAVY, "medium"), bottom=side(NAVY, "double"))


def _cell(ws, value, **kw):
    c = WriteOnlyCell(ws, value=value)
    for k, v in kw.items():
        setattr(c, k, v)
    return c


def _sheet_top(ws, st, width, title, subtitle):
    """Строки 1–4: оглавление, заголовок, подзаголовок, фирменная линия."""
    btn = [_cell(ws, None, fill=st.surface, border=Border(bottom=side())) for _ in range(width)]
    btn[0] = _cell(ws, "←  Оглавление", font=st.link, fill=st.surface, border=Border(bottom=side()))
    btn[0].hyperlink = "#'Оглавление'!A1"
    ws.append(btn)
    ws.append([_cell(ws, title.upper(), font=st.title)])
    ws.append([_cell(ws, subtitle, font=st.muted)])
    ws.append([_cell(ws, None, border=st.brand_line) for _ in range(width)])


def _table(ws, st, columns, rows, total=None):
    """columns: [(title, fmt)], rows: iterable of lists."""
    n = len(columns)
    ws.append([_cell(ws, t, font=st.hdr, fill=st.hdr_fill, alignment=st.center_wrap, border=st.hdr_border)
               for t, _ in columns])
    for i, row in enumerate(rows):
        zebra = i % 2 == 1
        out = []
        for j, ((_, fmt), v) in enumerate(zip(columns, row)):
            c = WriteOnlyCell(ws, value=v)
            c.font = st.body
            c.border = st.body_border if j < n - 1 else Border(bottom=side())
            if zebra:
                c.fill = st.zebra
            if fmt:
                c.number_format = fmt
                if fmt != FMT_DATE:
                    c.alignment = st.right
            out.append(c)
        ws.append(out)
    if total is not None:
        ws.append([_cell(ws, v, font=st.bold, fill=st.surface, border=st.total_border,
                         number_format=fmt or "General", alignment=st.right if fmt else st.left)
                   for (_, fmt), v in zip(columns, total)])


def _num(v):
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if np.isnan(f) else f


def _summary(df, key, label):
    g = (df.groupby(key, dropna=False)
         .agg(amount=("amount", "sum"), dt=("dt", "sum"), cr=("cr", "sum"),
              qty=("quant_dt", "sum"), qty_cr=("quant_cr", "sum"))
         .reset_index()
         .sort_values("amount", ascending=False))
    total = g["amount"].sum()
    g["share"] = g["amount"] / total if total else 0.0
    return g


def build_sales_excel_bytes(df: pd.DataFrame, start: str, end: str, scope: str) -> bytes:
    """Write-only: большие периоды пишутся быстро, оформление — построчно."""
    df = df.head(1_000_000).copy()
    for k in ("quant_dt", "quant_cr", "dt", "cr", "amount"):
        df[k] = pd.to_numeric(df[k], errors="coerce").fillna(0.0)
    df["date"] = pd.to_datetime(df["date"], errors="coerce")

    st = _Styles()
    period = f"{pd.to_datetime(start):%d.%m.%Y} — {pd.to_datetime(end):%d.%m.%Y}"
    params = f"Рубли и штуки · {period} · {scope} · выгрузка {datetime.now():%d.%m.%Y %H:%M}"
    wb = Workbook(write_only=True)

    # --- Оглавление
    toc = wb.create_sheet("Оглавление")
    toc.sheet_view.showGridLines = False
    for j, w in enumerate([3, 26, 3, 26, 3, 26, 3, 26], start=1):
        toc.column_dimensions[get_column_letter(j)].width = w
    toc.append([_cell(toc, "ПРОДАЖИ ЗА ПЕРИОД", font=st.title)])
    toc.append([_cell(toc, "Все строки продаж с группировками и итогами", font=Font(name=FONT, size=9, color=TEXT_2))])
    toc.append([_cell(toc, params, font=st.muted)])
    toc.append([_cell(toc, None, border=st.brand_line) for _ in range(8)])
    toc.append([])
    cards = [
        ("ВЫРУЧКА, ₽", df["amount"].sum(), FMT_MONEY),
        ("ПРОДАНО, ШТ.", df["quant_dt"].sum() - df["quant_cr"].sum(), FMT_QTY),
        ("ВОЗВРАТЫ, ₽", df["cr"].sum(), FMT_MONEY),
        ("СТРОК ПРОДАЖ", len(df), FMT_QTY),
    ]
    top = Border(top=side(NAVY, "medium"))
    toc.append(sum([[None, _cell(toc, l, font=st.card_label, border=top, fill=st.surface if i == 0 else PatternFill())]
                    for i, (l, _, _) in enumerate(cards)], []))
    toc.append(sum([[None, _cell(toc, float(v), font=st.card_value, number_format=f,
                                 alignment=st.left, fill=st.surface if i == 0 else PatternFill())]
                    for i, (_, v, f) in enumerate(cards)], []))
    toc.append([])
    toc.append([None, _cell(toc, "СОДЕРЖАНИЕ ОТЧЁТА", font=Font(name=FONT, size=12, bold=True, color=TEXT))])
    for name, desc in (
        ("Сводка", "Итоги по месяцам, магазинам, группам и категориям"),
        ("Продажи", "Все строки продаж: дата, товар, магазин, дизайнер, суммы"),
    ):
        link = _cell(toc, f"›  {name}", font=st.link10, fill=st.surface, border=Border(bottom=side()))
        link.hyperlink = f"#'{name}'!A1"
        toc.append([None, link, None, _cell(toc, desc, font=Font(name=FONT, size=9, color=TEXT_2))])

    # --- Сводка
    sm = wb.create_sheet("Сводка")
    sm.sheet_view.showGridLines = False
    for j, w in enumerate([30, 16, 16, 16, 12, 12, 12], start=1):
        sm.column_dimensions[get_column_letter(j)].width = w
    _sheet_top(sm, st, 7, "Сводка продаж", params)
    cols_sum = [("", None), ("Выручка, ₽", FMT_MONEY), ("Продажи, ₽", FMT_MONEY), ("Возвраты, ₽", FMT_MONEY),
                ("Продано, шт.", FMT_QTY), ("Возвраты, шт.", FMT_QTY), ("Доля выручки", FMT_SHARE)]
    df["_month"] = df["date"].dt.to_period("M")
    for key, label in (("_month", "Месяц"), ("store_group", "Магазин"), ("parent_cat", "Группа"), ("cat", "Категория")):
        g = _summary(df, key, label)
        if key == "_month":
            g = g.sort_values(key)
            g[key] = [f"{MONTHS[p.month - 1]} {p.year}" if pd.notna(p) else "—" for p in g[key]]
        sm.append([])
        sm.append([_cell(sm, label.upper(), font=Font(name=FONT, size=11, bold=True, color=NAVY_3))])
        cols = [(label, None)] + cols_sum[1:]
        _table(sm, st, cols,
               ([str(r[key]), _num(r["amount"]), _num(r["dt"]), _num(r["cr"]), _num(r["qty"]),
                 _num(r["qty_cr"]), _num(r["share"])] for _, r in g.iterrows()),
               total=["ИТОГО", g["amount"].sum(), g["dt"].sum(), g["cr"].sum(), g["qty"].sum(),
                      g["qty_cr"].sum(), 1.0])

    # --- Продажи
    ws = wb.create_sheet("Продажи")
    ws.sheet_view.showGridLines = False
    n = len(COLUMNS)
    for j, (_, _, _, w) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"{get_column_letter(FREEZE_AFTER + 1)}6"
    _sheet_top(ws, st, n, "Продажи за период", params)

    def rows():
        for rec in df[[k for k, _, _, _ in COLUMNS]].itertuples(index=False):
            out = []
            for (key, _, fmt, _), v in zip(COLUMNS, rec):
                if key == "date":
                    out.append(v.to_pydatetime() if pd.notna(v) else None)
                elif fmt:
                    out.append(_num(v))
                else:
                    out.append(None if v is None or (isinstance(v, float) and np.isnan(v)) else v)
            yield out

    total = ["ИТОГО"] + [None] * (n - 1)
    for j, (key, _, fmt, _) in enumerate(COLUMNS):
        if fmt and key != "date":
            total[j] = float(df[key].sum())
    _table(ws, st, [(t, f) for _, t, f, _ in COLUMNS], rows(), total=total)
    ws.auto_filter.ref = f"A5:{get_column_letter(n)}{len(df) + 5}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
