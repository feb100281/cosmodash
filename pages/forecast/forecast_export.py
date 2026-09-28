# pages/forecast/forecast_export.py
"""Excel-выгрузка плана продаж."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter

from pages.matrix.xl_brand import (
    FMT_DATE,
    FMT_MONEY,
    FMT_SHARE,
    NAVY,
    TOTAL_ROW,
    build_toc,
    fill,
    finalize,
    font,
    set_number_format,
    side,
    style_table_sheet,
)

MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
          "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]


def _clean(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    return v


def _total_row(ws, values, fmts):
    r = ws.max_row + 1
    for j, (v, fmt) in enumerate(zip(values, fmts), start=1):
        c = ws.cell(row=r, column=j, value=_clean(v))
        c.font = font(10, True)
        c.fill = fill(TOTAL_ROW)
        c.border = Border(top=side(NAVY, "medium"), bottom=side(NAVY, "double"))
        if fmt:
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right", vertical="center")


def build_forecast_excel_bytes(frame: pd.DataFrame, meta: dict) -> bytes:
    """
    Завершённые месяцы — только факт; текущий месяц — факт по последнюю дату
    продаж + план на остаток; будущие месяцы — только план.
    """
    df = frame.copy()
    df["ds"] = pd.to_datetime(df["ds"])
    cur = pd.to_datetime(meta["current_date"])
    horizon = pd.to_datetime(meta["horizon"])
    last_fact = pd.to_datetime(meta["last_fact_date"])
    df = df[df["ds"] <= horizon].sort_values("ds")

    is_fact = df["ds"] <= last_fact
    df["period"] = np.where(is_fact, "Факт", "План")
    df["fact"] = np.where(is_fact, pd.to_numeric(df["fact"], errors="coerce"), np.nan)
    for c in ("yhat", "yhat_lower", "yhat_upper"):
        df[c] = np.where(is_fact, np.nan, pd.to_numeric(df[c], errors="coerce"))

    params = (f"Рубли · факт по {last_fact:%d.%m.%Y} · план до {horizon:%d.%m.%Y} · "
              f"история с {pd.to_datetime(meta['history_start']):%d.%m.%Y}")

    wb = Workbook()
    wb.remove(wb.active)

    m = df.copy()
    m["year"] = m["ds"].dt.year
    m["month"] = m["ds"].dt.month
    monthly = (
        m.groupby(["year", "month"])
        .agg(fact=("fact", lambda s: s.sum(min_count=1)), plan=("yhat", lambda s: s.sum(min_count=1)))
        .reset_index()
    )
    monthly["total"] = monthly["fact"].fillna(0) + monthly["plan"].fillna(0)
    monthly["status"] = np.select(
        [monthly["plan"].isna(), monthly["fact"].isna()],
        ["Факт", "План"],
        default="Факт + план",
    )

    ws = wb.create_sheet("По месяцам")
    ws.append(["Год", "Месяц", "Статус", "Факт", "План", "Итого (факт + план)"])
    for r in monthly.itertuples():
        ws.append([int(r.year), MONTHS[int(r.month) - 1], r.status,
                   _clean(r.fact), _clean(r.plan), float(r.total)])
    hr = style_table_sheet(ws, title="План продаж по месяцам",
                           subtitle=params + " · завершённые месяцы — только факт",
                           header_row=1, freeze_col=4, key_headers=["Итого (факт + план)"])
    set_number_format(ws, hr, ["Факт", "План", "Итого (факт + план)"], FMT_MONEY)
    for r in range(hr + 1, ws.max_row + 1):
        st = ws.cell(row=r, column=3)
        if st.value == "План":
            st.font = font(10, True, NAVY)
        elif st.value == "Факт + план":
            st.font = font(10, True, "9A6100")
    _total_row(ws, ["ИТОГО", "", "", monthly["fact"].sum(min_count=1), monthly["plan"].sum(min_count=1),
                    monthly["total"].sum()], [None, None, None, FMT_MONEY, FMT_MONEY, FMT_MONEY])
    for j, w in enumerate([10, 14, 14, 18, 18, 22], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    yearly = monthly.groupby("year").agg(fact=("fact", lambda s: s.sum(min_count=1)),
                                          plan=("plan", lambda s: s.sum(min_count=1)),
                                          total=("total", "sum")).reset_index()
    ws = wb.create_sheet("По годам")
    ws.append(["Год", "Факт", "План", "Итого (факт + план)"])
    for r in yearly.itertuples():
        ws.append([int(r.year), _clean(r.fact), _clean(r.plan), float(r.total)])
    hr = style_table_sheet(ws, title="План продаж по годам", subtitle=params, header_row=1,
                           key_headers=["Итого (факт + план)"])
    set_number_format(ws, hr, ["Факт", "План", "Итого (факт + план)"], FMT_MONEY)
    for j, w in enumerate([10, 20, 20, 24], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # по дням
    ws = wb.create_sheet("По дням")
    ws.append(["Дата", "Тип", "Факт", "План", "План: нижняя граница", "План: верхняя граница"])
    for r in df[df["ds"] >= pd.Timestamp(year=last_fact.year - 1, month=1, day=1)].itertuples():
        ws.append([r.ds.to_pydatetime(), r.period, _clean(r.fact), _clean(r.yhat),
                   _clean(r.yhat_lower), _clean(r.yhat_upper)])
    hr = style_table_sheet(ws, title="План и факт по дням", subtitle=params, header_row=1)
    set_number_format(ws, hr, ["Факт", "План", "План: нижняя граница", "План: верхняя граница"], FMT_MONEY)
    for r in range(hr + 1, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = FMT_DATE
    for j, w in enumerate([12, 8, 16, 16, 20, 20], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # параметры модели
    ws = wb.create_sheet("Параметры")
    ws.append(["Параметр", "Значение"])
    rows = [
        ("Текущая дата", f"{cur:%d.%m.%Y}"),
        ("Последняя дата факта", f"{pd.to_datetime(meta['last_fact_date']):%d.%m.%Y}"),
        ("Горизонт планирования", f"{horizon:%d.%m.%Y}"),
        ("История с", f"{pd.to_datetime(meta['history_start']):%d.%m.%Y}"),
        ("Ошибка модели по месяцам",
         f"{meta['mape']:.1f}%".replace(".", ",") if meta.get("mape") is not None else "—"),
    ] + [(k, str(v)) for k, v in meta["params"].items()]
    plan_end = df.loc[df["period"] == "План", "ds"].max()
    rows.append(("План рассчитан по", f"{plan_end:%d.%m.%Y}" if pd.notna(plan_end) else "нет плановых дат"))
    for k, v in rows:
        ws.append([k, v])
    style_table_sheet(ws, title="Параметры расчёта", subtitle=params, header_row=1, freeze_col=1)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 26

    plan_total = float(df["yhat"].sum())
    year_cur = monthly[monthly["year"] == last_fact.year]
    build_toc(
        wb,
        title="План продаж",
        subtitle="Прогноз выручки (модель Prophet) до горизонта планирования",
        params=params,
        cards=[
            ("ПЛАН ДО ГОРИЗОНТА, ₽", plan_total, FMT_MONEY,
             f"{last_fact + pd.Timedelta(days=1):%d.%m.%Y} — {horizon:%d.%m.%Y}"),
            (f"ИТОГО {last_fact.year}, ₽", float(year_cur["total"].sum()), FMT_MONEY, "факт + план"),
            (f"ФАКТ {last_fact.year}, ₽", float(year_cur["fact"].sum()), FMT_MONEY, f"по {last_fact:%d.%m.%Y}"),
            ("ОШИБКА МОДЕЛИ", (meta.get("mape") or 0) / 100, FMT_SHARE, "по месяцам истории"),
        ],
        sheets=[
            ("По месяцам", "Завершённые месяцы — факт, текущий — факт + план, будущие — план"),
            ("По годам", "Итоги по годам"),
            ("По дням", "С начала прошлого года: факт по дням, дальше — план с диапазоном"),
            ("Параметры", "Даты и настройки модели, с которыми сделан расчёт"),
        ],
    )
    finalize(wb, ["Оглавление", "По месяцам", "По годам", "По дням", "Параметры"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
